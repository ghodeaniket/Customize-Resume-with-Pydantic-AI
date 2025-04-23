"""Document processing service for extracting text from various file formats.

This module provides functionality to extract text from PDF, DOCX, and TXT files.
"""

import io
import os
from enum import Enum
from pathlib import Path
from typing import List, Optional, Union

import docx
import fitz  # PyMuPDF
from fastapi import UploadFile
from PyPDF2 import PdfReader
from loguru import logger

from resume_customizer.core.config import settings
from resume_customizer.core.exceptions import DocumentProcessingError


class DocumentFormat(str, Enum):
    """Supported document formats."""
    PDF = "application/pdf"
    DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    TXT = "text/plain"
    UNKNOWN = "unknown"


class DocumentProcessor:
    """Service for processing document files and extracting text."""
    
    @staticmethod
    def get_format(file: UploadFile) -> DocumentFormat:
        """Determine the format of a document based on content type and filename.
        
        Args:
            file: The uploaded file
            
        Returns:
            DocumentFormat: The detected document format
        """
        # Check content type first
        content_type = file.content_type
        if content_type:
            for format_type in DocumentFormat:
                if content_type == format_type.value:
                    return format_type
        
        # Fall back to file extension
        if file.filename:
            extension = os.path.splitext(file.filename)[1].lower()
            if extension == '.pdf':
                return DocumentFormat.PDF
            elif extension == '.docx':
                return DocumentFormat.DOCX
            elif extension == '.txt':
                return DocumentFormat.TXT
        
        return DocumentFormat.UNKNOWN
    
    @staticmethod
    async def is_valid_file(file: UploadFile) -> bool:
        """Check if a file is valid (correct format and size).
        
        Args:
            file: The uploaded file
            
        Returns:
            bool: True if the file is valid, False otherwise
        """
        # Check file size
        file.file.seek(0, os.SEEK_END)
        size = file.file.tell()
        file.file.seek(0)  # Reset file position
        
        if size > settings.MAX_UPLOAD_SIZE:
            logger.warning(f"File too large: {size} bytes (max: {settings.MAX_UPLOAD_SIZE})")
            return False
        
        # Check file format
        format_type = DocumentProcessor.get_format(file)
        if format_type == DocumentFormat.UNKNOWN:
            logger.warning(f"Unknown file format: {file.content_type} / {file.filename}")
            return False
        
        return True
    
    @staticmethod
    async def save_file(file: UploadFile) -> Path:
        """Save an uploaded file to disk.
        
        Args:
            file: The uploaded file
            
        Returns:
            Path: The path to the saved file
            
        Raises:
            DocumentProcessingError: If there's an error saving the file
        """
        # Ensure upload directory exists
        os.makedirs(settings.UPLOAD_DIRECTORY, exist_ok=True)
        
        try:
            # Create a safe filename
            filename = Path(file.filename).name if file.filename else "unnamed_file"
            file_path = Path(settings.UPLOAD_DIRECTORY) / filename
            
            # Save the file
            contents = await file.read()
            with open(file_path, "wb") as f:
                f.write(contents)
            
            logger.info(f"Saved file to {file_path}")
            return file_path
            
        except Exception as e:
            error_msg = f"Failed to save file: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg)
        finally:
            await file.seek(0)  # Reset file position for further processing
    
    @staticmethod
    async def extract_text_from_pdf(file_content: bytes) -> str:
        """Extract text from a PDF file using PyMuPDF (more reliable than PyPDF2).
        
        Args:
            file_content: The raw PDF file content
            
        Returns:
            str: The extracted text
            
        Raises:
            DocumentProcessingError: If there's an error extracting text
        """
        try:
            # First try with PyMuPDF for better text extraction
            text = ""
            pdf_file = io.BytesIO(file_content)
            
            try:
                pdf_document = fitz.open(stream=pdf_file, filetype="pdf")
                for page_num in range(len(pdf_document)):
                    page = pdf_document.load_page(page_num)
                    text += page.get_text()
                pdf_document.close()
            except Exception as mupdf_error:
                logger.warning(f"PyMuPDF extraction failed, falling back to PyPDF2: {str(mupdf_error)}")
                
                # Fall back to PyPDF2 if PyMuPDF fails
                pdf_file.seek(0)  # Reset file position
                reader = PdfReader(pdf_file)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            
            if not text.strip():
                logger.warning("Extracted empty text from PDF")
                raise ValueError("No text could be extracted from the PDF file")
            
            logger.info(f"Successfully extracted {len(text)} characters from PDF")
            return text
            
        except Exception as e:
            error_msg = f"Failed to extract text from PDF: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg, document_type="PDF")
    
    @staticmethod
    async def extract_text_from_docx(file_content: bytes) -> str:
        """Extract text from a DOCX file.
        
        Args:
            file_content: The raw DOCX file content
            
        Returns:
            str: The extracted text
            
        Raises:
            DocumentProcessingError: If there's an error extracting text
        """
        try:
            docx_file = io.BytesIO(file_content)
            doc = docx.Document(docx_file)
            
            # Extract text from paragraphs
            paragraphs = [para.text for para in doc.paragraphs]
            
            # Extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.append(cell.text)
            
            text = "\n".join(filter(None, paragraphs))
            
            if not text.strip():
                logger.warning("Extracted empty text from DOCX")
                raise ValueError("No text could be extracted from the DOCX file")
            
            logger.info(f"Successfully extracted {len(text)} characters from DOCX")
            return text
            
        except Exception as e:
            error_msg = f"Failed to extract text from DOCX: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg, document_type="DOCX")
    
    @staticmethod
    async def extract_text_from_txt(file_content: bytes) -> str:
        """Extract text from a TXT file.
        
        Args:
            file_content: The raw TXT file content
            
        Returns:
            str: The extracted text
            
        Raises:
            DocumentProcessingError: If there's an error extracting text
        """
        try:
            # Try different encodings
            encodings = ['utf-8', 'latin1', 'cp1252']
            text = None
            
            for encoding in encodings:
                try:
                    text = file_content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue
            
            if text is None:
                raise ValueError("Could not decode text file with any of the attempted encodings")
            
            if not text.strip():
                logger.warning("Extracted empty text from TXT")
                raise ValueError("The text file is empty")
            
            logger.info(f"Successfully extracted {len(text)} characters from TXT")
            return text
            
        except Exception as e:
            error_msg = f"Failed to extract text from TXT: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg, document_type="TXT")
    
    @staticmethod
    async def extract_text(
        file: Union[UploadFile, bytes, str, Path],
        format_type: Optional[DocumentFormat] = None
    ) -> str:
        """Extract text from a document file.
        
        This method supports various input types:
        - UploadFile: A FastAPI uploaded file
        - bytes: Raw file content
        - str: File path or raw text content
        - Path: File path
        
        Args:
            file: The document file to process
            format_type: The document format (will be auto-detected if None)
            
        Returns:
            str: The extracted text
            
        Raises:
            DocumentProcessingError: If there's an error extracting text
        """
        start_time = logger.time.time()
        logger.info(f"Starting text extraction from document")
        
        try:
            # Process based on input type
            if isinstance(file, UploadFile):
                # FastAPI UploadFile
                format_type = format_type or DocumentProcessor.get_format(file)
                content = await file.read()
                await file.seek(0)  # Reset file position
                
            elif isinstance(file, bytes):
                # Raw file content
                content = file
                if format_type is None:
                    raise ValueError("format_type must be provided when file is bytes")
                
            elif isinstance(file, (str, Path)):
                # File path or raw text
                if isinstance(file, str) and os.path.exists(file) or isinstance(file, Path):
                    # It's a file path
                    with open(file, "rb") as f:
                        content = f.read()
                    
                    # Auto-detect format from extension if not provided
                    if format_type is None:
                        ext = os.path.splitext(str(file))[1].lower()
                        if ext == '.pdf':
                            format_type = DocumentFormat.PDF
                        elif ext == '.docx':
                            format_type = DocumentFormat.DOCX
                        elif ext == '.txt':
                            format_type = DocumentFormat.TXT
                        else:
                            raise ValueError(f"Unsupported file extension: {ext}")
                else:
                    # It's raw text content
                    return file if isinstance(file, str) else file.read_text()
            else:
                raise TypeError(f"Unsupported file type: {type(file)}")
            
            # Extract text based on format
            if format_type == DocumentFormat.PDF:
                text = await DocumentProcessor.extract_text_from_pdf(content)
            elif format_type == DocumentFormat.DOCX:
                text = await DocumentProcessor.extract_text_from_docx(content)
            elif format_type == DocumentFormat.TXT:
                text = await DocumentProcessor.extract_text_from_txt(content)
            else:
                raise ValueError(f"Unsupported document format: {format_type}")
            
            elapsed_time = logger.time.time() - start_time
            logger.info(f"Text extraction completed in {elapsed_time:.2f} seconds")
            
            return text
            
        except DocumentProcessingError:
            # Re-raise existing DocumentProcessingError
            raise
        except Exception as e:
            error_msg = f"Failed to extract text from document: {str(e)}"
            logger.error(error_msg)
            raise DocumentProcessingError(error_msg)
